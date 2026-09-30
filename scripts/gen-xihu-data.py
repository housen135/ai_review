#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
把「西湖英才评审数据源.xlsx」转成 src/data/xihuProjects.ts。

用法(在项目根目录):
    python scripts/gen-xihu-data.py [xlsx路径]

默认读根目录的 西湖英才评审数据源.xlsx。改了表就重跑这个脚本,不要手改生成的 .ts。

依赖:openpyxl(pip install openpyxl)
"""
import json
import os
import sys

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else '西湖英才评审数据源.xlsx'
OUT = os.path.join('src', 'data', 'xihuProjects.ts')

# 列序对应表头:PPT申报人 / PPT来源文件 / 评审组别 / 分类组别 / 项目名称 / 项目简介 /
#              毕业院校 / 拟创办企业名称 / 专家1-7评分 / 总分 / 平均分 / 专家意见
COL_APPLICANT, COL_REVIEW_GROUP, COL_GROUP = 0, 2, 3
COL_NAME, COL_INTRO, COL_SCHOOL, COL_COMPANY = 4, 5, 6, 7
COL_EXPERT0, COL_TOTAL, COL_AVG, COL_COMMENT = 8, 15, 16, 17
EXPERT_COUNT = 7


def text(row, i):
    v = row[i] if i < len(row) else None
    return '' if v is None else str(v).strip()


def number(row, i):
    t = text(row, i)
    if not t:
        return None
    try:
        return round(float(t), 2)
    except ValueError:
        return None


def js(value):
    # json.dumps 产出的就是合法的 JS 字面量;再挡一下 JS 里非法的行分隔符
    return (
        json.dumps(value, ensure_ascii=False)
        .replace(' ', '\\u2028')
        .replace(' ', '\\u2029')
    )


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))[1:]

    items, missing_scores = [], []
    for row in rows:
        name = text(row, COL_NAME)
        if not name:
            continue
        scores = [number(row, COL_EXPERT0 + k) for k in range(EXPERT_COUNT)]
        if any(s is None for s in scores):
            missing_scores.append(name)
        items.append({
            'id': 'XH%03d' % (len(items) + 1),
            'name': name,
            'applicant': text(row, COL_APPLICANT),
            'group': text(row, COL_GROUP),
            'reviewGroup': text(row, COL_REVIEW_GROUP),
            'school': text(row, COL_SCHOOL),
            'company': text(row, COL_COMPANY),
            'intro': text(row, COL_INTRO),
            'expertScores': scores,
            'totalScore': number(row, COL_TOTAL),
            'avgScore': number(row, COL_AVG),
            'expertComment': text(row, COL_COMMENT),
        })

    lines = [
        '/**',
        ' * 第二十三批西湖英才项目数据。',
        ' *',
        ' * 本文件由 scripts/gen-xihu-data.py 从「西湖英才评审数据源.xlsx」生成,请勿手改 ——',
        ' * 要更新数据请改表后重跑:python scripts/gen-xihu-data.py',
        ' * 原始列名对照:PPT申报人 / PPT来源文件 / 评审组别 / 分类组别 / 项目名称 / 项目简介 /',
        ' *                毕业院校 / 拟创办企业名称 / 专家1-7评分 / 总分 / 平均分 / 专家意见',
        ' */',
        '',
        'export interface XihuProject {',
        '  id: string;',
        '  name: string;',
        '  /** PPT申报人 */',
        '  applicant: string;',
        '  /** 分类组别,对应概览页的八个项目分组 */',
        '  group: string;',
        '  /** 评审组别,如「人工智能1组」 */',
        '  reviewGroup: string;',
        '  school: string;',
        '  company: string;',
        '  intro: string;',
        '  /** 七位专家的评分,缺分的位为 null */',
        '  expertScores: (number | null)[];',
        '  totalScore: number | null;',
        '  avgScore: number | null;',
        '  /** 专家组综合意见(优势/劣势) */',
        '  expertComment: string;',
        '}',
        '',
        'export const XIHU_PROJECTS: XihuProject[] = [',
    ]
    for it in items:
        lines.append('  { ' + ', '.join(f'{k}: {js(v)}' for k, v in it.items()) + ' },')
    lines += ['];', '']

    with open(OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))

    print('条目 %d -> %s (%.0f KB)' % (len(items), OUT, os.path.getsize(OUT) / 1024))
    if missing_scores:
        print('注意:以下项目有专家缺分,已按 null 保留 —— ' + '、'.join(missing_scores))


if __name__ == '__main__':
    main()
